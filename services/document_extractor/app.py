from __future__ import annotations

import csv
import io
import os
from typing import Any

from fastapi import FastAPI, File, Header, HTTPException, UploadFile
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

app = FastAPI(title="UDC Document Extractor", version="1.0.0")


class ExtractionResponse(BaseModel):
    text: str
    page_count: int | None = None
    extraction_method: str
    warnings: list[str]
    metadata: dict[str, Any]


def require_token(authorization: str | None) -> None:
    expected = os.environ.get("DOCUMENT_EXTRACTOR_TOKEN", "").strip()
    if not expected:
        raise HTTPException(status_code=503, detail="extractor_not_configured")
    if authorization != f"Bearer {expected}":
        raise HTTPException(status_code=401, detail="unauthorized")


def extract_pdf(data: bytes) -> ExtractionResponse:
    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    warnings: list[str] = []

    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
            warnings.append(f"page_{index}_text_extraction_failed")
        pages.append(f"\n--- PAGE {index} ---\n{text.strip()}")

    joined = "\n".join(pages).strip()
    if len(joined.replace("\n", " ").strip()) < max(100, len(reader.pages) * 30):
        warnings.append("low_text_density_possible_scan")

    return ExtractionResponse(
        text=joined,
        page_count=len(reader.pages),
        extraction_method="pypdf",
        warnings=warnings,
        metadata={
            "pdf_metadata": {
                str(k): str(v)
                for k, v in (reader.metadata or {}).items()
                if v is not None
            }
        },
    )


def extract_docx(data: bytes) -> ExtractionResponse:
    doc = Document(io.BytesIO(data))
    chunks: list[str] = []

    for paragraph in doc.paragraphs:
        value = paragraph.text.strip()
        if value:
            chunks.append(value)

    for table_index, table in enumerate(doc.tables, start=1):
        chunks.append(f"\n--- TABLE {table_index} ---")
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))

    return ExtractionResponse(
        text="\n".join(chunks).strip(),
        extraction_method="python-docx",
        warnings=[],
        metadata={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)},
    )


def extract_xlsx(data: bytes) -> ExtractionResponse:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks: list[str] = []

    for sheet in workbook.worksheets:
        chunks.append(f"\n--- SHEET: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                chunks.append(" | ".join(values))

    return ExtractionResponse(
        text="\n".join(chunks).strip(),
        extraction_method="openpyxl",
        warnings=[],
        metadata={"sheets": workbook.sheetnames},
    )


def extract_csv(data: bytes) -> ExtractionResponse:
    decoded = None
    encoding = None
    for candidate in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            decoded = data.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue

    if decoded is None:
        raise ValueError("csv_decode_failed")

    reader = csv.reader(io.StringIO(decoded))
    lines = [" | ".join(row) for row in reader]
    return ExtractionResponse(
        text="\n".join(lines).strip(),
        extraction_method="python-csv",
        warnings=[],
        metadata={"encoding": encoding},
    )


def extract_text(data: bytes) -> ExtractionResponse:
    for candidate in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return ExtractionResponse(
                text=data.decode(candidate),
                extraction_method="plain-text",
                warnings=[],
                metadata={"encoding": candidate},
            )
        except UnicodeDecodeError:
            continue
    raise ValueError("text_decode_failed")


@app.get("/health")
def health() -> dict[str, bool]:
    return {"ok": True}


@app.post("/extract", response_model=ExtractionResponse)
async def extract(
    file: UploadFile = File(...),
    authorization: str | None = Header(default=None),
) -> ExtractionResponse:
    require_token(authorization)

    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="empty_file")
    if len(data) > 30 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="file_too_large")

    name = (file.filename or "").lower()
    mime = (file.content_type or "").lower()

    try:
        if mime == "application/pdf" or name.endswith(".pdf"):
            return extract_pdf(data)
        if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or name.endswith(".docx"):
            return extract_docx(data)
        if mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" or name.endswith(".xlsx"):
            return extract_xlsx(data)
        if mime in ("text/csv", "application/csv") or name.endswith(".csv"):
            return extract_csv(data)
        if mime.startswith("text/") or name.endswith((".txt", ".md")):
            return extract_text(data)

        raise HTTPException(status_code=415, detail="unsupported_document_type")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail={"code": "extraction_failed", "type": type(exc).__name__},
        ) from exc
