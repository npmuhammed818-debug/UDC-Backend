from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

import app as extractor


client = TestClient(extractor.app)
TOKEN = "test-extractor-token"


def _upload(filename: str, data: bytes, content_type: str = "application/octet-stream", token: str = TOKEN):
    return client.post(
        "/extract",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": (filename, data, content_type)},
    )


def test_extract_requires_configured_bearer_token(monkeypatch):
    monkeypatch.delenv("DOCUMENT_EXTRACTOR_TOKEN", raising=False)
    unconfigured = client.post(
        "/extract",
        files={"file": ("requirement.txt", b"Copper, 100 MT", "text/plain")},
    )
    assert unconfigured.status_code == 503
    assert unconfigured.json()["detail"] == "extractor_not_configured"

    monkeypatch.setenv("DOCUMENT_EXTRACTOR_TOKEN", TOKEN)
    unauthorized = _upload("requirement.txt", b"Copper, 100 MT", token="wrong-token")
    assert unauthorized.status_code == 401
    assert unauthorized.json()["detail"] == "unauthorized"

    accepted = _upload("requirement.txt", b"Copper, 100 MT", "text/plain")
    assert accepted.status_code == 200
    assert accepted.json()["text"] == "Copper, 100 MT"
    assert accepted.json()["extraction_method"] == "plain-text"


def test_extract_rejects_empty_unsupported_and_oversized_files(monkeypatch):
    monkeypatch.setenv("DOCUMENT_EXTRACTOR_TOKEN", TOKEN)

    empty = _upload("empty.txt", b"", "text/plain")
    assert empty.status_code == 400
    assert empty.json()["detail"] == "empty_file"

    unsupported = _upload("offer.bin", b"binary", "application/octet-stream")
    assert unsupported.status_code == 415
    assert unsupported.json()["detail"] == "unsupported_document_type"

    oversized = _upload("large.txt", b"x" * (30 * 1024 * 1024 + 1), "text/plain")
    assert oversized.status_code == 413
    assert oversized.json()["detail"] == "file_too_large"


def test_extract_csv_docx_and_xlsx_content(monkeypatch):
    monkeypatch.setenv("DOCUMENT_EXTRACTOR_TOKEN", TOKEN)

    csv_result = _upload(
        "offer.csv",
        b"\xef\xbb\xbfproduct,quantity\ncopper,100 MT",
        "text/csv",
    )
    assert csv_result.status_code == 200
    assert csv_result.json()["text"] == "product | quantity\ncopper | 100 MT"
    assert csv_result.json()["metadata"]["encoding"] == "utf-8-sig"

    document = Document()
    document.add_paragraph("Copper Millberry")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Quantity"
    table.cell(0, 1).text = "100 MT"
    docx_bytes = BytesIO()
    document.save(docx_bytes)

    docx_result = _upload(
        "offer.docx",
        docx_bytes.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert docx_result.status_code == 200
    assert "Copper Millberry" in docx_result.json()["text"]
    assert "Quantity | 100 MT" in docx_result.json()["text"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offer"
    sheet.append(["Product", "Quantity"])
    sheet.append(["Copper", "100 MT"])
    xlsx_bytes = BytesIO()
    workbook.save(xlsx_bytes)

    xlsx_result = _upload(
        "offer.xlsx",
        xlsx_bytes.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert xlsx_result.status_code == 200
    assert "SHEET: Offer" in xlsx_result.json()["text"]
    assert "Copper | 100 MT" in xlsx_result.json()["text"]
